"""评估报告生成器：按指定模拟时间区间汇总指标/评分/路口/事件/图表。

支持三种输出格式：
  - md   ：Markdown（图片以 base64 data URI 内嵌，任何浏览器可直接预览）
  - docx ：Word 文档（python-docx）
  - pdf  ：PDF（reportlab + CID 中文字体）

约定：报告文本不使用任何 emoji，纯文本排版；图表由 matplotlib 绘制（Agg 无头模式）。

说明：matplotlib / python-docx / reportlab 采用**延迟导入**——即使环境缺少
报告依赖，后端其余功能（仿真/算法/智能体）照常运行，仅报告接口返回缺依赖提示。
"""

import base64
import io
import os
import tempfile
from datetime import datetime

# ── 中文字体（Windows） ─────────────────────────────────────
_CN_FONT = False
_PLT = None  # matplotlib.pyplot（延迟加载）


def _setup_font() -> bool:
    """初始化 matplotlib 中文字体。缺 matplotlib 时返回 False（图表接口将报缺依赖）。"""
    global _CN_FONT, _PLT
    if _CN_FONT:
        return True
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib import font_manager
    except ImportError:
        return False
    for path in (r"C:\Windows\Fonts\simhei.ttf",
                 r"C:\Windows\Fonts\msyh.ttc",
                 r"C:\Windows\Fonts\simsun.ttc"):
        if not os.path.isfile(path):
            continue
        try:
            font_manager.fontManager.addfont(path)
            name = font_manager.FontProperties(fname=path).get_name()
            plt.rcParams["font.sans-serif"] = [name, "DejaVu Sans"]
            plt.rcParams["axes.unicode_minus"] = False
            _CN_FONT = True
            _PLT = plt
            return True
        except Exception:  # noqa: BLE001 ttc 集合可能不受支持
            continue
    _PLT = plt
    return True


def _fmt_num(v, digits: int = 1) -> str:
    try:
        return f"{float(v):.{digits}f}"
    except (TypeError, ValueError):
        return "—"


def _series_stats(rows: list[dict]) -> dict:
    """[{step, value}] → 均值/最低/最高/首/尾。"""
    vals = [r["value"] for r in rows]
    if not vals:
        return {"count": 0, "mean": None, "min": None, "max": None,
                "first": None, "last": None}
    return {
        "count": len(vals),
        "mean": round(sum(vals) / len(vals), 3),
        "min": round(min(vals), 3),
        "max": round(max(vals), 3),
        "first": round(vals[0], 3),
        "last": round(vals[-1], 3),
    }


def los_level(delay: float) -> str:
    if delay <= 10:
        return "A"
    if delay <= 20:
        return "B"
    if delay <= 35:
        return "C"
    if delay <= 55:
        return "D"
    if delay <= 80:
        return "E"
    return "F"


class ReportGenerator:
    """报告生成器。collect 收集区间数据，charts 绘图，render_* 输出各格式。"""

    def __init__(self, runtime):
        self.rt = runtime

    # ── 数据收集 ────────────────────────────────────────────

    def collect(self, start: int, end: int) -> dict:
        rt = self.rt
        data = {
            "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "net_file": rt._net_path or "",
            "net_summary": rt.network_summary() or {},
            "scheme": rt.status().get("scheme", "none"),
            "range": {"start": max(0, int(start)), "end": max(0, int(end))},
            "hist": {},
            "realtime": {},
            "spotlight": {},
            "score": None,
            "events": [],
        }
        try:
            data["realtime"] = rt.realtime_metrics()
        except Exception:  # noqa: BLE001
            pass
        try:
            data["spotlight"] = rt.spotlight()
        except Exception:  # noqa: BLE001
            pass
        try:
            data["score"] = rt.evaluate_score()
        except Exception:  # noqa: BLE001
            pass
        try:
            evs = rt.list_events()
            lo, hi = data["range"]["start"], data["range"]["end"]
            data["events"] = [e for e in evs if lo <= int(e.get("step", 0)) <= hi]
        except Exception:  # noqa: BLE001
            pass
        for metric in ("avg_speed", "avg_delay", "throughput", "queue_length",
                       "fuel", "co2"):
            try:
                rows = rt.store.query(metric, data["range"]["start"],
                                      data["range"]["end"], 1, "overall")
            except Exception:  # noqa: BLE001
                rows = []
            data["hist"][metric] = rows
        return data

    # ── 图表（返回 图名 → PNG bytes） ────────────────────────

    def charts(self, data: dict) -> dict[str, bytes]:
        """生成图表。缺 matplotlib 时抛出带安装提示的错误。"""
        global _PLT
        if _PLT is None:
            if not _setup_font():
                raise RuntimeError("缺少 matplotlib，无法生成图表。"
                                   "请执行: pip install matplotlib")
        plt = _PLT
        out: dict[str, bytes] = {}
        hist = data["hist"]

        def fig_bytes(fig) -> bytes:
            buf = io.BytesIO()
            fig.savefig(buf, format="png", dpi=130, bbox_inches="tight")
            plt.close(fig)
            return buf.getvalue()

        steps = [r["step"] for r in hist.get("avg_speed", [])]
        if steps:
            fig, ax = plt.subplots(figsize=(8, 3))
            ax.plot(steps, [r["value"] * 3.6 for r in hist["avg_speed"]],
                    color="#2f9e44", linewidth=1.5)
            ax.set_title("平均速度（km/h）" if _CN_FONT else "Average Speed (km/h)")
            ax.set_xlabel("仿真步" if _CN_FONT else "Simulation Step")
            ax.set_ylabel("km/h")
            ax.grid(True, linestyle="--", alpha=0.4)
            fig.tight_layout()
            out["speed"] = fig_bytes(fig)

        qsteps = [r["step"] for r in hist.get("queue_length", [])]
        if qsteps:
            fig, ax = plt.subplots(figsize=(8, 3))
            ax.plot(qsteps, [r["value"] for r in hist["queue_length"]],
                    color="#e8590c", linewidth=1.5)
            ax.set_title("平均排队（辆/边）" if _CN_FONT else "Avg Queue (veh/edge)")
            ax.set_xlabel("仿真步" if _CN_FONT else "Simulation Step")
            ax.set_ylabel("辆" if _CN_FONT else "veh")
            ax.grid(True, linestyle="--", alpha=0.4)
            fig.tight_layout()
            out["queue"] = fig_bytes(fig)

        tsteps = [r["step"] for r in hist.get("throughput", [])]
        if tsteps:
            fig, ax = plt.subplots(figsize=(8, 3))
            ax.plot(tsteps, [r["value"] for r in hist["throughput"]],
                    color="#1971c2", linewidth=1.5)
            ax.set_title("累计到达车辆" if _CN_FONT else "Cumulative Arrivals")
            ax.set_xlabel("仿真步" if _CN_FONT else "Simulation Step")
            ax.set_ylabel("辆" if _CN_FONT else "veh")
            ax.grid(True, linestyle="--", alpha=0.4)
            fig.tight_layout()
            out["throughput"] = fig_bytes(fig)

        inters = data.get("realtime", {}).get("intersections", {})
        if inters:
            items = sorted(inters.items(), key=lambda kv: kv[1].get("waiting_time", 0),
                           reverse=True)
            ids = [k for k, _ in items[:12]]
            waits = [v.get("waiting_time", 0) for _, v in items[:12]]
            fig, ax = plt.subplots(figsize=(8, 3.4))
            ax.bar(ids, waits, color="#7048e8")
            ax.set_title("各路口平均等待（s）" if _CN_FONT else "Intersection Avg Wait (s)")
            ax.set_xlabel("路口 id" if _CN_FONT else "Intersection")
            ax.set_ylabel("s")
            ax.tick_params(axis="x", rotation=45, labelsize=8)
            ax.grid(True, axis="y", linestyle="--", alpha=0.4)
            fig.tight_layout()
            out["tls_wait"] = fig_bytes(fig)
        return out

    # ── Markdown ────────────────────────────────────────────

    def render_md(self, data: dict, charts: dict[str, bytes],
                  sections: list[str]) -> str:
        L: list[str] = []
        lo, hi = data["range"]["start"], data["range"]["end"]
        L.append("# 交通仿真评估报告")
        L.append("")
        L.append("- 生成时间：" + data["generated_at"])
        L.append("- 路网文件：" + (data["net_file"] or "—"))
        summ = data["net_summary"]
        if summ:
            L.append("- 路网概况：" + ("，".join(f"{k} {v}" for k, v in summ.items())
                                      if isinstance(summ, dict) else str(summ)))
        L.append("- 控制方案：" + str(data["scheme"]))
        L.append(f"- 评估区间：第 {lo} ~ 第 {hi} 仿真步（模拟秒），共 {max(0, hi - lo)} 步")
        L.append("")
        L.append("> 说明：除历史统计外，指标为结束时刻的实时近似快照；精确评分数据可在平台的全屏数据面板查看。")
        L.append("")

        if "overview" in sections:
            L.append("## 1 全局概况（结束时刻）")
            L.append("")
            L.append("| 指标 | 数值 |")
            L.append("| --- | --- |")
            ov = data["realtime"].get("overall", {})
            rows = [
                ("在网车辆", str(ov.get("vehicle_count", 0)) + " 辆"),
                ("平均速度", _fmt_num((ov.get("avg_speed", 0) or 0) * 3.6) + " km/h"),
                ("平均延误", _fmt_num(ov.get("avg_delay", 0)) + " s"),
                ("平均等待", _fmt_num(ov.get("avg_waiting_time", 0)) + " s"),
                ("平均排队", _fmt_num(ov.get("avg_queue_length", 0), 2) + " 辆/边"),
                ("累计到达", str(ov.get("total_throughput", 0)) + " 辆"),
            ]
            sp = data.get("spotlight", {})
            rows.append(("完成率", (f"{sp.get('completion_rate', 0) * 100:.1f}%"
                                    if sp.get("total_departed") else "—")))
            rows.append(("累计出发", str(sp.get("total_departed", 0)) + " 辆"))
            for k, v in rows:
                L.append(f"| {k} | {v} |")
            L.append("")

        if "score" in sections and data["score"]:
            sc = data["score"]
            L.append("## 2 综合评分（老评价平台分层聚合公式，实时近似）")
            L.append("")
            if sc.get("passed"):
                L.append(f"- 综合评分：**{sc['score']:.3f}**（满分 1.0，越大越好；"
                         "公式 = 效率^0.5 × 稳定^0.3 × 公平^0.2）")
                L.append("")
                L.append("| 子项 | 得分 | 满分 | 方向 |")
                L.append("| --- | --- | --- | --- |")
                for k in ("efficiency", "stability", "fairness"):
                    v = sc.get(k, {}).get("value", 0)
                    L.append(f"| {k} | {_fmt_num(v, 3)} | 1.0 | 越大越好 |")
                L.append("")
            else:
                L.append("硬性筛选未通过：")
                for r in sc.get("reasons", []):
                    L.append(f"- {r}")
                L.append("")

        if "history" in sections:
            L.append("## 3 区间历史统计")
            L.append("")
            L.append("| 指标 | 采样数 | 均值 | 最低 | 最高 |")
            L.append("| --- | --- | --- | --- | --- |")
            labels = {
                "avg_speed": "平均速度 (m/s)",
                "avg_delay": "平均延误 (s)",
                "throughput": "累计到达 (辆)",
                "queue_length": "平均排队 (辆/边)",
                "fuel": "油耗 (L)",
                "co2": "CO2 (g)",
            }
            for metric, label in labels.items():
                st = _series_stats(data["hist"].get(metric, []))
                if st["count"] == 0:
                    continue
                L.append(f"| {label} | {st['count']} | {_fmt_num(st['mean'], 2)} | "
                         f"{_fmt_num(st['min'], 2)} | {_fmt_num(st['max'], 2)} |")
            L.append("")

        if "intersections" in sections:
            L.append("## 4 各路口评估（结束时刻）")
            L.append("")
            inters = data["realtime"].get("intersections", {})
            if inters:
                L.append("| 路口 | 排队(辆) | 平均等待(s) | LOS |")
                L.append("| --- | --- | --- | --- |")
                for tid, v in sorted(inters.items()):
                    wait = v.get("waiting_time", 0) or 0
                    L.append(f"| {tid} | {v.get('queue_length', 0)} | "
                             f"{_fmt_num(wait)} | {los_level(wait)} |")
            else:
                L.append("（无路口数据）")
            L.append("")

        if "charts" in sections and charts:
            L.append("## 5 图表")
            L.append("")
            order = [("speed", "5.1 平均速度曲线"),
                     ("queue", "5.2 平均排队曲线"),
                     ("throughput", "5.3 累计到达曲线"),
                     ("tls_wait", "5.4 各路口平均等待")] 
            for key, caption in order:
                if key not in charts:
                    continue
                b64 = base64.b64encode(charts[key]).decode("ascii")
                L.append(f"### {caption}")
                L.append("")
                L.append(f"![{caption}](data:image/png;base64,{b64})")
                L.append("")

        if "events" in sections:
            L.append("## 6 事件日志（区间内）")
            L.append("")
            if data["events"]:
                L.append("| 类型 | 仿真步 | 受影响边 |")
                L.append("| --- | --- | --- |")
                for e in data["events"][-30:]:
                    edges = e.get("params", {}).get("edge_ids", [])
                    L.append(f"| {e.get('event_type', '—')} | {e.get('step', 0)} | "
                             f"{'、'.join(map(str, edges)) if edges else '全入口'} |")
            else:
                L.append("（区间内无事件）")
            L.append("")
        return "\n".join(L)

    # ── Word (docx) ─────────────────────────────────────────

    def render_docx(self, data: dict, charts: dict[str, bytes],
                    sections: list[str]) -> bytes:
        try:
            from docx import Document
            from docx.shared import Pt, Inches, RGBColor
        except ImportError as exc:
            raise RuntimeError("缺少 python-docx，无法生成 Word 报告。"
                               "请执行: pip install python-docx") from exc

        doc = Document()
        lo, hi = data["range"]["start"], data["range"]["end"]

        def h(text, level=1):
            head = doc.add_heading(text, level=level)
            for run in head.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
            return head

        def para(text, bold=False):
            p = doc.add_paragraph()
            r = p.add_run(text)
            r.bold = bold
            return p

        h("交通仿真评估报告", 0)
        para(f"生成时间：{data['generated_at']}")
        para(f"路网文件：{data['net_file'] or '—'}")
        para(f"控制方案：{data['scheme']}")
        para(f"评估区间：第 {lo} ~ 第 {hi} 仿真步，共 {max(0, hi - lo)} 步")
        para("说明：除历史统计外，指标为结束时刻的实时近似快照。")

        if "overview" in sections:
            h("1 全局概况（结束时刻）", 1)
            ov = data["realtime"].get("overall", {})
            sp = data.get("spotlight", {})
            comp = (f"{sp.get('completion_rate', 0) * 100:.1f}%"
                    if sp.get("total_departed") else "—")
            for k, v in [
                ("在网车辆", f"{ov.get('vehicle_count', 0)} 辆"),
                ("平均速度", f"{_fmt_num((ov.get('avg_speed', 0) or 0) * 3.6)} km/h"),
                ("平均延误", f"{_fmt_num(ov.get('avg_delay', 0))} s"),
                ("平均等待", f"{_fmt_num(ov.get('avg_waiting_time', 0))} s"),
                ("平均排队", f"{_fmt_num(ov.get('avg_queue_length', 0), 2)} 辆/边"),
                ("累计到达", f"{ov.get('total_throughput', 0)} 辆"),
                ("完成率", comp),
                ("累计出发", f"{sp.get('total_departed', 0)} 辆"),
            ]:
                para(f"{k}：{v}")

        if "score" in sections and data["score"]:
            h("2 综合评分（实时近似）", 1)
            sc = data["score"]
            if sc.get("passed"):
                para(f"综合评分：{sc['score']:.3f} / 1.0（越大越好，"
                     "效率^0.5 × 稳定^0.3 × 公平^0.2）")
                table = doc.add_table(rows=1, cols=4)
                table.style = "Light Grid Accent 1"
                for i, c in enumerate(("子项", "得分", "满分", "方向")):
                    table.rows[0].cells[i].text = c
                for k in ("efficiency", "stability", "fairness"):
                    row = table.add_row()
                    row.cells[0].text = k
                    row.cells[1].text = _fmt_num(sc.get(k, {}).get("value", 0), 3)
                    row.cells[2].text = "1.0"
                    row.cells[3].text = "越大越好"
            else:
                para("硬性筛选未通过：")
                for r in sc.get("reasons", []):
                    para(f"  - {r}")

        if "history" in sections:
            h("3 区间历史统计", 1)
            table = doc.add_table(rows=1, cols=5)
            table.style = "Light Grid Accent 1"
            for i, c in enumerate(("指标", "采样数", "均值", "最低", "最高")):
                table.rows[0].cells[i].text = c
            labels = {
                "avg_speed": "平均速度 (m/s)", "avg_delay": "平均延误 (s)",
                "throughput": "累计到达 (辆)", "queue_length": "平均排队 (辆/边)",
                "fuel": "油耗 (L)", "co2": "CO2 (g)",
            }
            for metric, label in labels.items():
                st = _series_stats(data["hist"].get(metric, []))
                if st["count"] == 0:
                    continue
                row = table.add_row()
                row.cells[0].text = label
                row.cells[1].text = str(st["count"])
                row.cells[2].text = _fmt_num(st["mean"], 2)
                row.cells[3].text = _fmt_num(st["min"], 2)
                row.cells[4].text = _fmt_num(st["max"], 2)

        if "intersections" in sections:
            h("4 各路口评估（结束时刻）", 1)
            inters = data["realtime"].get("intersections", {})
            if inters:
                table = doc.add_table(rows=1, cols=4)
                table.style = "Light Grid Accent 1"
                for i, c in enumerate(("路口", "排队(辆)", "平均等待(s)", "LOS")):
                    table.rows[0].cells[i].text = c
                for tid, v in sorted(inters.items()):
                    wait = v.get("waiting_time", 0) or 0
                    row = table.add_row()
                    row.cells[0].text = str(tid)
                    row.cells[1].text = str(v.get("queue_length", 0))
                    row.cells[2].text = _fmt_num(wait)
                    row.cells[3].text = los_level(wait)
            else:
                para("（无路口数据）")

        if "charts" in sections:
            order = [("speed", "平均速度曲线"), ("queue", "平均排队曲线"),
                     ("throughput", "累计到达曲线"), ("tls_wait", "各路口平均等待")]
            for key, caption in order:
                if key not in charts:
                    continue
                h(f"5 {caption}", 2)
                doc.add_picture(io.BytesIO(charts[key]), width=Inches(6))

        if "events" in sections:
            h("6 事件日志（区间内）", 1)
            if data["events"]:
                table = doc.add_table(rows=1, cols=3)
                table.style = "Light Grid Accent 1"
                for i, c in enumerate(("类型", "仿真步", "受影响边")):
                    table.rows[0].cells[i].text = c
                for e in data["events"][-30:]:
                    edges = e.get("params", {}).get("edge_ids", [])
                    row = table.add_row()
                    row.cells[0].text = str(e.get("event_type", "—"))
                    row.cells[1].text = str(e.get("step", 0))
                    row.cells[2].text = "、".join(map(str, edges)) if edges else "全入口"
            else:
                para("（区间内无事件）")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    # ── PDF ─────────────────────────────────────────────────

    def render_pdf(self, data: dict, charts: dict[str, bytes],
                   sections: list[str]) -> bytes:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                            Table, TableStyle, Image)
            from reportlab.lib import colors
        except ImportError as exc:
            raise RuntimeError("缺少 reportlab，无法生成 PDF 报告。"
                               "请执行: pip install reportlab") from exc

        try:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            _font = "STSong-Light"
        except Exception:  # noqa: BLE001
            _font = "Helvetica"

        def st(size=10, leading=14, bold=False, title=False):
            return ParagraphStyle(
                f"s{size}{bold}", fontName=_font, fontSize=size,
                leading=leading, spaceAfter=4,
                textColor=colors.HexColor("#111111"))

        story = []
        H1 = st(16, 22, True)
        H2 = st(12, 16, True)
        BODY = st(10)
        lo, hi = data["range"]["start"], data["range"]["end"]

        story.append(Paragraph("交通仿真评估报告", H1))
        story.append(Spacer(1, 6))
        for line in (
            f"生成时间：{data['generated_at']}",
            f"路网文件：{data['net_file'] or '—'}",
            f"控制方案：{data['scheme']}",
            f"评估区间：第 {lo} ~ 第 {hi} 仿真步，共 {max(0, hi - lo)} 步",
            "说明：除历史统计外，指标为结束时刻的实时近似快照。",
        ):
            story.append(Paragraph(line, BODY))
        story.append(Spacer(1, 8))

        def table(data_rows, widths=None):
            t = Table(data_rows, colWidths=widths, hAlign="LEFT")
            t.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), _font),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9ecef")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#ced4da")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            return t

        if "overview" in sections:
            story.append(Paragraph("1 全局概况（结束时刻）", H2))
            ov = data["realtime"].get("overall", {})
            sp = data.get("spotlight", {})
            comp = (f"{sp.get('completion_rate', 0) * 100:.1f}%"
                    if sp.get("total_departed") else "—")
            rows = [["指标", "数值"],
                    ["在网车辆", f"{ov.get('vehicle_count', 0)} 辆"],
                    ["平均速度", f"{_fmt_num((ov.get('avg_speed', 0) or 0) * 3.6)} km/h"],
                    ["平均延误", f"{_fmt_num(ov.get('avg_delay', 0))} s"],
                    ["平均等待", f"{_fmt_num(ov.get('avg_waiting_time', 0))} s"],
                    ["平均排队", f"{_fmt_num(ov.get('avg_queue_length', 0), 2)} 辆/边"],
                    ["累计到达", f"{ov.get('total_throughput', 0)} 辆"],
                    ["完成率", comp],
                    ["累计出发", f"{sp.get('total_departed', 0)} 辆"]]
            story.append(table(rows, widths=[1.6 * inch, 2.6 * inch]))
            story.append(Spacer(1, 8))

        if "score" in sections and data["score"]:
            story.append(Paragraph("2 综合评分（实时近似）", H2))
            sc = data["score"]
            if sc.get("passed"):
                story.append(Paragraph(
                    f"综合评分：{sc['score']:.3f} / 1.0（越大越好，"
                    "效率^0.5 × 稳定^0.3 × 公平^0.2）", BODY))
                rows = [["子项", "得分", "满分", "方向"]]
                for k in ("efficiency", "stability", "fairness"):
                    rows.append([k, _fmt_num(sc.get(k, {}).get("value", 0), 3),
                                 "1.0", "越大越好"])
                story.append(table(rows, widths=[1.4 * inch] * 4))
            else:
                story.append(Paragraph("硬性筛选未通过：" + "；".join(
                    sc.get("reasons", [])), BODY))
            story.append(Spacer(1, 8))

        if "history" in sections:
            story.append(Paragraph("3 区间历史统计", H2))
            rows = [["指标", "采样数", "均值", "最低", "最高"]]
            labels = {
                "avg_speed": "平均速度 (m/s)", "avg_delay": "平均延误 (s)",
                "throughput": "累计到达 (辆)", "queue_length": "平均排队 (辆/边)",
                "fuel": "油耗 (L)", "co2": "CO2 (g)",
            }
            for metric, label in labels.items():
                s = _series_stats(data["hist"].get(metric, []))
                if s["count"] == 0:
                    continue
                rows.append([label, str(s["count"]), _fmt_num(s["mean"], 2),
                             _fmt_num(s["min"], 2), _fmt_num(s["max"], 2)])
            story.append(table(rows, widths=[1.7 * inch, 0.7 * inch, 1.0 * inch,
                                             1.0 * inch, 1.0 * inch]))
            story.append(Spacer(1, 8))

        if "intersections" in sections:
            story.append(Paragraph("4 各路口评估（结束时刻）", H2))
            inters = data["realtime"].get("intersections", {})
            if inters:
                rows = [["路口", "排队(辆)", "平均等待(s)", "LOS"]]
                for tid, v in sorted(inters.items()):
                    wait = v.get("waiting_time", 0) or 0
                    rows.append([str(tid), str(v.get("queue_length", 0)),
                                 _fmt_num(wait), los_level(wait)])
                story.append(table(rows, widths=[1.2 * inch] * 4))
            else:
                story.append(Paragraph("（无路口数据）", BODY))
            story.append(Spacer(1, 8))

        if "charts" in sections:
            order = [("speed", "5 平均速度曲线"), ("queue", "6 平均排队曲线"),
                     ("throughput", "7 累计到达曲线"), ("tls_wait", "8 各路口平均等待")]
            for key, caption in order:
                if key not in charts:
                    continue
                story.append(Paragraph(caption, H2))
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
                    f.write(charts[key])
                    tmp = f.name
                story.append(Image(tmp, width=6.4 * inch, height=2.6 * inch))
                story.append(Spacer(1, 6))

        if "events" in sections:
            story.append(Paragraph("9 事件日志（区间内）", H2))
            if data["events"]:
                rows = [["类型", "仿真步", "受影响边"]]
                for e in data["events"][-30:]:
                    edges = e.get("params", {}).get("edge_ids", [])
                    rows.append([str(e.get("event_type", "—")),
                                 str(e.get("step", 0)),
                                 "、".join(map(str, edges)) if edges else "全入口"])
                story.append(table(rows, widths=[1.2 * inch, 0.9 * inch, 2.7 * inch]))
            else:
                story.append(Paragraph("（区间内无事件）", BODY))

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4,
                                leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                                topMargin=0.7 * inch, bottomMargin=0.7 * inch)
        doc.build(story)
        return buf.getvalue()

    # ── 入口 ────────────────────────────────────────────────

    def generate(self, fmt: str, start: int, end: int,
                 sections: list[str]) -> tuple[str, bytes]:
        """生成报告。返回 (文件名, 内容字节)。fmt ∈ {md, docx, pdf}。"""
        data = self.collect(start, end)
        charts = self.charts(data) if "charts" in sections else {}
        if fmt == "md":
            content = self.render_md(data, charts, sections).encode("utf-8")
        elif fmt == "docx":
            content = self.render_docx(data, charts, sections)
        elif fmt == "pdf":
            content = self.render_pdf(data, charts, sections)
        else:
            raise ValueError(f"不支持的格式: {fmt}")
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"评估报告_{ts}.{fmt}", content
