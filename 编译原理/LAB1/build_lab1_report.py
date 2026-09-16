from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"D:\大三作业\编译原理\LAB1")
SOURCE = ROOT / "实验报告_C--词法与语法分析.md"
OUTPUT = ROOT / "实验报告_C--词法与语法分析.docx"


def set_run_font(run, name="宋体", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, "Arial", 9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def make_doc():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for style_name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    add_page_number(section.footer.paragraphs[0])

    i = 0
    first = True
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(18)
            r = p.add_run(line[2:].strip())
            set_run_font(r, "微软雅黑", 21, True)
            first = False
            i += 1
            continue

        if first and not line.startswith("#"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(line.strip())
            set_run_font(r, "宋体", 11, False, "404040")
            i += 1
            continue

        if line.startswith("## "):
            p = doc.add_paragraph(line[3:].strip(), style="Heading 1")
            p.paragraph_format.first_line_indent = Cm(0)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(line[4:].strip(), style="Heading 2")
            p.paragraph_format.first_line_indent = Cm(0)
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                if not set(lines[i].replace("|", "").replace("-", "").replace(":", "").strip()):
                    i += 1
                    continue
                table_lines.append([x.strip() for x in lines[i].strip().strip("|").split("|")])
                i += 1
            if table_lines:
                cols = max(len(row) for row in table_lines)
                table = doc.add_table(rows=len(table_lines), cols=cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True
                for ri, row in enumerate(table_lines):
                    for ci in range(cols):
                        value = row[ci] if ci < len(row) else ""
                        set_cell_text(table.cell(ri, ci), value, ri == 0, "FFFFFF" if ri == 0 else "000000")
                        set_cell_borders(table.cell(ri, ci))
                        if ri == 0:
                            shade_cell(table.cell(ri, ci), "365F91")
                        elif ri % 2 == 0:
                            shade_cell(table.cell(ri, ci), "F3F6FA")
                doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue

        if line.startswith("    "):
            code_lines = []
            while i < len(lines) and (lines[i].startswith("    ") or not lines[i].strip()):
                code_lines.append(lines[i][4:] if lines[i].startswith("    ") else "")
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.right_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run("\n".join(code_lines).rstrip())
            set_run_font(r, "Consolas", 9, False, "202020")
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.74)
            r = p.add_run(line[2:].strip())
            set_run_font(r)
            i += 1
            continue

        if len(line) >= 3 and line[0].isdigit() and line[1:3] == ". ":
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(line[3:].strip())
            set_run_font(r)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        r = p.add_run(line.strip())
        set_run_font(r)
        i += 1

    doc.core_properties.title = "C--词法分析与语法分析实验报告"
    doc.core_properties.subject = "编译原理 Lab 1"
    doc.core_properties.author = "学生"
    doc.save(OUTPUT)


if __name__ == "__main__":
    make_doc()
