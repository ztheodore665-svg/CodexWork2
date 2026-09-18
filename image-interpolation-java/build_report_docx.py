from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parent
OUTPUT = PROJECT / "实验报告_代码与结果.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="宋体", size=12, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(paragraph, first_line=True, before=0, after=6, line=1.35):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Cm(0.74)


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    format_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, name="黑体", size=15 if level == 1 else 13, bold=True)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_cell_shading(cell, "44546A")
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, size=10.5, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            if row_index % 2 == 1:
                set_cell_shading(cell, "F2F5F7")
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 1 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=10.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    return table


def add_code_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.right_indent = Cm(0.8)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=9.5)
    return p


def add_result_grid(doc, scale, files):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"cameraman {scale} 倍结果")
    set_run_font(r, name="黑体", size=11, bold=True)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    captions = ["最近邻插值", "双线性插值", "双三次插值"]
    for i, (path, caption) in enumerate(zip(files, captions)):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_cell_borders(cell, color="FFFFFF", size="0")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p_img = cell.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(path), width=Inches(2.05))
        p_cap = cell.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(3)
        r = p_cap.add_run(caption)
        set_run_font(r, size=9.5)
        cell.width = Inches(2.15)
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(12)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("数字图像处理实验报告")
    set_run_font(footer_run, size=9, color="666666")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("数字图像处理实验报告")
    set_run_font(title_run, name="黑体", size=20, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run("Java 图像插值代码文件与 src 图片处理结果")
    set_run_font(subtitle_run, name="黑体", size=13)

    info = doc.add_table(rows=2, cols=4)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info_data = [["课程", "数字图像处理（SSE317）", "姓名", "__________"],
                 ["实验内容", "三种图像插值与实际图片缩放", "学号", "__________"]]
    for r_i, row in enumerate(info.rows):
        for c_i, cell in enumerate(row.cells):
            cell.text = ""
            set_cell_borders(cell)
            set_cell_margins(cell, top=110, bottom=110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(info_data[r_i][c_i])
            set_run_font(run, size=10.5, bold=(c_i % 2 == 0))
            if c_i % 2 == 0:
                set_cell_shading(cell, "EAF0F6")
    widths = [0.75, 2.55, 0.75, 1.15]
    for row in info.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)

    add_heading(doc, "1 代码文件说明", 1)
    add_body(doc, "本项目使用 Java 标准库完成图像读取、像素访问、插值计算和 PNG 输出。代码文件之间职责清晰：ImageInterpolator.java 负责算法，Main.java 负责单张图片，BatchMain.java 负责处理 src 文件夹中的全部图片，InterpolationTest.java 负责基础正确性验证。")
    add_table(doc,
              ["文件", "主要作用"],
              [
                  ["ImageInterpolator.java", "核心算法文件。统一完成坐标映射、边界钳位，并实现最近邻、双线性和双三次三种插值。"],
                  ["InterpolationMethod.java", "定义 NEAREST、BILINEAR、BICUBIC 三种插值方法及命令行名称。"],
                  ["Main.java", "单张图片处理入口。读取输入图片，默认生成 0.5 倍和 3 倍的六张结果图。"],
                  ["BatchMain.java", "批处理入口。扫描 src 文件夹中的 PNG、JPG、BMP、GIF、TIF 和 TIFF 图片，并按算法、倍率分类保存。"],
                  ["InterpolationTest.java", "基础测试文件。验证边界像素、输出尺寸和双线性中心像素，运行后输出测试是否通过。"],
                  ["run.ps1 / batch.ps1", "PowerShell 辅助脚本，自动编译 Java 源文件并执行单张图片或批处理任务。"],
                  ["package.ps1", "根据学号和姓名生成作业提交 ZIP 压缩包。"],
              ], widths=[2.0, 4.25])

    add_heading(doc, "2 程序运行方式", 1)
    add_body(doc, "本次实验直接使用 src 文件夹中的 19 张图片进行批处理。进入 image-interpolation-java 文件夹后运行以下命令：", bold_prefix=None)
    add_code_line(doc, ".\\batch.ps1")
    add_body(doc, "程序将结果输出到 results-src 文件夹，目录结构如下：")
    add_code_line(doc, "results-src/nearest/0.5x/     最近邻缩小结果")
    add_code_line(doc, "results-src/nearest/3x/       最近邻放大结果")
    add_code_line(doc, "results-src/bilinear/0.5x/    双线性缩小结果")
    add_code_line(doc, "results-src/bilinear/3x/      双线性放大结果")
    add_code_line(doc, "results-src/bicubic/0.5x/     双三次缩小结果")
    add_code_line(doc, "results-src/bicubic/3x/       双三次放大结果")

    add_heading(doc, "3 src 图片处理结果", 1)
    add_body(doc, "src 文件夹中共有 19 张图片，包含 cameraman、house、jetplane、lake、lena、mandril、peppers、pirate、walkbridge 和人物图像等。程序对每张图片分别执行三种插值，并生成 0.5 倍和 3 倍两种尺寸。")
    add_table(doc,
              ["项目", "结果"],
              [
                  ["输入图片数量", "19 张"],
                  ["插值方法", "最近邻、双线性、双三次"],
                  ["缩放倍率", "0.5 倍、3 倍"],
                  ["输出图片总数", "19 × 3 × 2 = 114 张"],
                  ["输出格式", "PNG"],
              ], widths=[2.0, 4.25])
    add_body(doc, "以 cameraman.tif 为例，原图尺寸为 512×512，三种方法的 0.5 倍结果均为 256×256，3 倍结果均为 1536×1536。对于 256×256 的输入图像，输出尺寸对应为 128×128 和 768×768。")

    add_heading(doc, "4 代表性结果对比", 1)
    nearest_half = PROJECT / "results-src" / "nearest" / "0.5x" / "cameraman.png"
    bilinear_half = PROJECT / "results-src" / "bilinear" / "0.5x" / "cameraman.png"
    bicubic_half = PROJECT / "results-src" / "bicubic" / "0.5x" / "cameraman.png"
    add_result_grid(doc, "0.5", [nearest_half, bilinear_half, bicubic_half])
    nearest_three = PROJECT / "results-src" / "nearest" / "3x" / "cameraman.png"
    bilinear_three = PROJECT / "results-src" / "bilinear" / "3x" / "cameraman.png"
    bicubic_three = PROJECT / "results-src" / "bicubic" / "3x" / "cameraman.png"
    add_result_grid(doc, "3", [nearest_three, bilinear_three, bicubic_three])
    add_body(doc, "从 0.5 倍结果可以看出，最近邻保留了离散像素特征，但细节跳变较明显；双线性结果更加平滑；双三次在轮廓过渡方面更自然。从 3 倍结果可以看出，最近邻会出现明显的方块和阶梯边缘，双线性能够减弱块状感但略显柔和，双三次的轮廓连续性和视觉清晰度通常更好。")

    add_heading(doc, "5 结果文件位置", 1)
    add_body(doc, "完整的 114 张处理结果保存在项目目录的 results-src 文件夹中。每个结果文件均以原图片文件名命名，并转换为 PNG 格式，便于在 Word、图片查看器或实验报告中直接使用。")
    add_code_line(doc, "image-interpolation-java/results-src/")

    add_heading(doc, "6 实验小结", 1)
    add_body(doc, "本次实验完成了 Java 插值程序对 src 文件夹全部图片的批量处理。代码文件分别承担算法实现、单图处理、批量处理和测试职责，生成结果按照算法和倍率分目录保存，便于逐项比较。实际结果显示，最近邻适合追求速度的场景，双线性适合一般图像缩放，双三次适合更重视放大后边缘和纹理质量的场景。")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
